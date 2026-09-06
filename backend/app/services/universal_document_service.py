"""Universal ASKODOX file/document extraction foundation."""
from __future__ import annotations

import csv
import io
import json
from pathlib import Path
from typing import Any, Dict

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader


class UniversalDocumentService:
    MAX_TEXT_CHARS = 120_000

    def analyze(self, *, file_bytes: bytes, filename: str, mime_type: str | None = None) -> Dict[str, Any]:
        if not file_bytes:
            raise ValueError("empty file")

        name = Path(filename or "upload").name
        suffix = Path(name).suffix.casefold()
        mime = str(mime_type or "").casefold().strip()

        if suffix == ".pdf" or mime == "application/pdf":
            text, metadata = self._pdf(file_bytes)
            kind = "pdf"
        elif suffix in {".xlsx", ".xlsm"} or "spreadsheet" in mime or "excel" in mime:
            text, metadata = self._xlsx(file_bytes)
            kind = "spreadsheet"
        elif suffix == ".docx" or "wordprocessingml" in mime:
            text, metadata = self._docx(file_bytes)
            kind = "document"
        elif suffix == ".csv" or mime in {"text/csv", "application/csv"}:
            text, metadata = self._csv(file_bytes)
            kind = "csv"
        elif suffix == ".json" or mime == "application/json":
            text, metadata = self._json(file_bytes)
            kind = "json"
        elif suffix in {".txt", ".md", ".log"} or mime.startswith("text/"):
            text = self._decode(file_bytes)
            metadata = {"lines": len(text.splitlines())}
            kind = "text"
        else:
            raise ValueError("unsupported file type")

        normalized = "\n".join(line.rstrip() for line in text.splitlines()).strip()
        return {
            "kind": kind,
            "filename": name,
            "mime_type": mime_type or None,
            "text": normalized[: self.MAX_TEXT_CHARS],
            "truncated": len(normalized) > self.MAX_TEXT_CHARS,
            "metadata": metadata,
        }

    @staticmethod
    def _decode(data: bytes) -> str:
        for encoding in ("utf-8-sig", "utf-8", "utf-16", "latin-1"):
            try:
                return data.decode(encoding)
            except UnicodeDecodeError:
                continue
        return data.decode("utf-8", errors="replace")

    def _pdf(self, data: bytes) -> tuple[str, Dict[str, Any]]:
        reader = PdfReader(io.BytesIO(data))
        pages = []
        for index, page in enumerate(reader.pages, start=1):
            pages.append(f"[Page {index}]\n{page.extract_text() or ''}")
        return "\n\n".join(pages), {"pages": len(reader.pages)}

    def _xlsx(self, data: bytes) -> tuple[str, Dict[str, Any]]:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        blocks = []
        sheet_rows: Dict[str, int] = {}
        try:
            for sheet in workbook.worksheets:
                rows = []
                count = 0
                for row in sheet.iter_rows(values_only=True):
                    values = ["" if value is None else str(value) for value in row]
                    if any(value.strip() for value in values):
                        rows.append("\t".join(values).rstrip())
                        count += 1
                sheet_rows[sheet.title] = count
                blocks.append(f"[Sheet: {sheet.title}]\n" + "\n".join(rows))
        finally:
            workbook.close()
        return "\n\n".join(blocks), {"sheets": list(sheet_rows), "rows_by_sheet": sheet_rows}

    def _docx(self, data: bytes) -> tuple[str, Dict[str, Any]]:
        document = Document(io.BytesIO(data))
        paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
        tables = []
        for table_index, table in enumerate(document.tables, start=1):
            lines = []
            for row in table.rows:
                lines.append("\t".join(cell.text.strip() for cell in row.cells))
            tables.append(f"[Table {table_index}]\n" + "\n".join(lines))
        text = "\n".join(paragraphs + tables)
        return text, {"paragraphs": len(paragraphs), "tables": len(document.tables)}

    def _csv(self, data: bytes) -> tuple[str, Dict[str, Any]]:
        decoded = self._decode(data)
        rows = list(csv.reader(io.StringIO(decoded)))
        text = "\n".join("\t".join(cell for cell in row) for row in rows)
        columns = max((len(row) for row in rows), default=0)
        return text, {"rows": len(rows), "columns": columns}

    def _json(self, data: bytes) -> tuple[str, Dict[str, Any]]:
        parsed = json.loads(self._decode(data))
        text = json.dumps(parsed, ensure_ascii=False, indent=2)
        if isinstance(parsed, dict):
            shape = "object"
            items = len(parsed)
        elif isinstance(parsed, list):
            shape = "array"
            items = len(parsed)
        else:
            shape = type(parsed).__name__
            items = 1
        return text, {"shape": shape, "items": items}
